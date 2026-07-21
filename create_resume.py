from io import BytesIO
import os
from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Inches, Pt, RGBColor

from reportlab.lib import colors
from reportlab.lib.pagesizes import letter
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import inch
from reportlab.platypus import Image, Paragraph, SimpleDocTemplate, Spacer


def generate_resume_docx(data, template=1):
    """Generate resume in DOCX format with selected template."""
    doc = Document()

    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.5)
        section.bottom_margin = Inches(0.5)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    if template == 1:
        _create_template1_docx(doc, data)
    elif template == 3:
        _create_template3_docx(doc, data)
    else:
        _create_template2_docx(doc, data)

    buffer = BytesIO()
    doc.save(buffer)
    buffer.seek(0)
    return buffer


BASE_DIR = Path(__file__).resolve().parent


def _resolve_photo_path(photo_value):
    if not photo_value:
        return None

    if isinstance(photo_value, Path):
        photo_value = str(photo_value)

    if not isinstance(photo_value, str):
        return None

    photo_value = photo_value.strip()
    if not photo_value:
        return None

    candidate_paths = []
    if os.path.isabs(photo_value):
        candidate_paths.append(photo_value)
    else:
        candidate_paths.extend([
            photo_value,
            str(BASE_DIR / photo_value),
            str(BASE_DIR / "static" / photo_value),
            str(BASE_DIR / "static" / "uploads" / photo_value),
            str(BASE_DIR / "uploads" / photo_value),
        ])

    for candidate in candidate_paths:
        if os.path.exists(candidate):
            return candidate

    return None


def _add_photo(doc, data):
    if data.get("show_photo", True):
        photo_path = _resolve_photo_path(data.get("photo"))
        if photo_path:
            try:
                doc.add_picture(photo_path, width=Inches(1.25))
                last_paragraph = doc.paragraphs[-1]
                last_paragraph.alignment = WD_ALIGN_PARAGRAPH.CENTER
            except Exception:
                pass


def _create_template1_docx(doc, data):
    """Classic template with a polished profile section."""
    _add_photo(doc, data)

    name_para = doc.add_paragraph()
    name_run = name_para.add_run(data.get("name", "Your Name"))
    name_run.font.size = Pt(24)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(79, 108, 255)
    name_para.alignment = WD_ALIGN_PARAGRAPH.CENTER

    role_para = doc.add_paragraph(data.get("role", "Professional"))
    role_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    role_run = role_para.runs[0]
    role_run.font.size = Pt(12)
    role_run.font.color.rgb = RGBColor(118, 75, 162)

    contact_para = doc.add_paragraph()
    contact_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
    contact_run = contact_para.add_run(f"📧 {data.get('email', '')} | 📱 {data.get('phone', '')}")
    contact_run.font.size = Pt(10)

    extra_links = []
    if data.get("linkedin"):
        extra_links.append(f"LinkedIn: {data.get('linkedin')}")
    if data.get("portfolio"):
        extra_links.append(f"Portfolio: {data.get('portfolio')}")
    if extra_links:
        links_para = doc.add_paragraph(" | ".join(extra_links))
        links_para.alignment = WD_ALIGN_PARAGRAPH.CENTER
        links_para.runs[0].font.size = Pt(9)

    doc.add_paragraph()

    for heading, value in [
        ("Professional Summary", data.get("summary") or data.get("goal", "")),
        ("Certifications", data.get("certifications", "")),
        ("Projects", data.get("projects", "")),
        ("Languages", data.get("languages", "")),
        ("Achievements", data.get("achievements", "")),
        ("Technical Skills", data.get("skills", "")),
        ("Professional Experience", data.get("experience", "")),
        ("Education", data.get("education", "")),
    ]:
        if not value:
            continue
        heading_para = doc.add_heading(heading, level=2)
        heading_para.runs[0].font.color.rgb = RGBColor(79, 108, 255)
        heading_para.runs[0].font.size = Pt(14)
        doc.add_paragraph(value)
        doc.add_paragraph()


def _create_template2_docx(doc, data):
    """Modern template with tighter spacing and a minimal layout."""
    name_para = doc.add_paragraph()
    name_run = name_para.add_run(data.get("name", "Your Name"))
    name_run.font.size = Pt(22)
    name_run.font.bold = True
    name_run.font.color.rgb = RGBColor(118, 75, 162)

    contact_para = doc.add_paragraph()
    contact_run = contact_para.add_run(f"{data.get('email', '')} | {data.get('phone', '')} | {data.get('role', '')}")
    contact_run.font.size = Pt(10)
    contact_run.font.color.rgb = RGBColor(79, 108, 255)

    extra_links = []
    if data.get("linkedin"):
        extra_links.append(f"LinkedIn: {data.get('linkedin')}")
    if data.get("portfolio"):
        extra_links.append(f"Portfolio: {data.get('portfolio')}")
    if extra_links:
        links_para = doc.add_paragraph(" | ".join(extra_links))
        links_para.runs[0].font.size = Pt(9)

    doc.add_paragraph()

    for heading, value in [
        ("Professional Summary", data.get("summary") or data.get("goal", "")),
        ("Certifications", data.get("certifications", "")),
        ("Projects", data.get("projects", "")),
        ("Languages", data.get("languages", "")),
        ("Achievements", data.get("achievements", "")),
        ("Professional Experience", data.get("experience", "")),
        ("Education", data.get("education", "")),
        ("Technical Skills", data.get("skills", "")),
    ]:
        if not value:
            continue
        heading_para = doc.add_heading(heading, level=2)
        heading_para.runs[0].font.color.rgb = RGBColor(118, 75, 162)
        heading_para.runs[0].font.size = Pt(14)
        doc.add_paragraph(value)
        doc.add_paragraph()


def _create_template3_docx(doc, data):
    """Creative template with larger header and callout sections."""
    _add_photo(doc, data)

    title_para = doc.add_paragraph()
    title_run = title_para.add_run(data.get("name", "Your Name"))
    title_run.font.size = Pt(26)
    title_run.font.bold = True
    title_run.font.color.rgb = RGBColor(31, 41, 55)

    role_para = doc.add_paragraph(data.get("role", "Professional"))
    role_para.alignment = WD_ALIGN_PARAGRAPH.LEFT
    role_run = role_para.runs[0]
    role_run.font.size = Pt(12)
    role_run.font.italic = True
    role_run.font.color.rgb = RGBColor(79, 108, 255)

    contact_para = doc.add_paragraph()
    contact_run = contact_para.add_run(f"{data.get('email', '')} • {data.get('phone', '')}")
    contact_run.font.size = Pt(10)

    extra_links = []
    if data.get("linkedin"):
        extra_links.append(f"LinkedIn: {data.get('linkedin')}")
    if data.get("portfolio"):
        extra_links.append(f"Portfolio: {data.get('portfolio')}")
    if extra_links:
        links_para = doc.add_paragraph(" | ".join(extra_links))
        links_para.runs[0].font.size = Pt(9)

    doc.add_paragraph()

    summary_para = doc.add_paragraph(data.get("summary") or data.get("goal", ""))
    summary_para.style = "Intense Quote"

    for heading, value in [
        ("Certifications", data.get("certifications", "")),
        ("Projects", data.get("projects", "")),
        ("Languages", data.get("languages", "")),
        ("Achievements", data.get("achievements", "")),
        ("Technical Skills", data.get("skills", "")),
        ("Professional Experience", data.get("experience", "")),
        ("Education", data.get("education", "")),
    ]:
        if not value:
            continue
        heading_para = doc.add_heading(heading, level=2)
        heading_para.runs[0].font.color.rgb = RGBColor(79, 108, 255)
        heading_para.runs[0].font.size = Pt(14)
        doc.add_paragraph(value)
        doc.add_paragraph()


def _accent_color(value):
    accent_map = {
        "blue": colors.HexColor("#4f6cff"),
        "purple": colors.HexColor("#7c3aed"),
        "green": colors.HexColor("#059669"),
        "orange": colors.HexColor("#ea580c"),
    }
    return accent_map.get(value, colors.HexColor("#4f6cff"))


def generate_resume_pdf(data, template=1):
    """Generate resume in PDF format with selected template and optional profile image."""
    buffer = BytesIO()
    doc = SimpleDocTemplate(
        buffer,
        pagesize=letter,
        rightMargin=0.75 * inch,
        leftMargin=0.75 * inch,
        topMargin=0.5 * inch,
        bottomMargin=0.5 * inch,
    )

    story = []
    styles = getSampleStyleSheet()
    accent = _accent_color(data.get("accent_color", "blue"))

    title_style = ParagraphStyle(
        "CustomTitle",
        parent=styles["Normal"],
        fontSize=24,
        textColor=accent,
        spaceAfter=6,
        alignment=1,
        fontName="Helvetica-Bold",
    )

    heading_style = ParagraphStyle(
        "CustomHeading",
        parent=styles["Heading2"],
        fontSize=14,
        textColor=accent,
        spaceAfter=10,
        spaceBefore=10,
        fontName="Helvetica-Bold",
        borderBottom=True,
        borderBottomColor=accent,
        borderBottomWidth=1,
    )

    contact_style = ParagraphStyle(
        "Contact",
        parent=styles["Normal"],
        fontSize=10,
        textColor=colors.HexColor("#4b5563"),
        alignment=1,
        spaceAfter=12,
    )

    body_style = ParagraphStyle(
        "Body",
        parent=styles["Normal"],
        fontSize=10,
        spaceAfter=8,
        leading=12,
    )

    if template == 3:
        story.append(Paragraph(data.get("name", "Your Name"), title_style))
        story.append(Paragraph(data.get("role", "Professional"), ParagraphStyle(
            "Role",
            parent=body_style,
            fontName="Helvetica-Oblique",
            textColor=accent,
            alignment=1,
            spaceAfter=10,
        )))
        story.append(Paragraph(f"{data.get('email', '')} • {data.get('phone', '')}", contact_style))
    else:
        story.append(Paragraph(data.get("name", "Your Name"), title_style))
        story.append(Paragraph(f"{data.get('email', '')} | {data.get('phone', '')}", contact_style))

    if data.get("show_photo", True):
        photo_path = _resolve_photo_path(data.get("photo"))
        if photo_path:
            story.append(Spacer(1, 0.1 * inch))
            story.append(Image(photo_path, width=1.2 * inch, height=1.4 * inch, kind="proportional", hAlign="CENTER"))
            story.append(Spacer(1, 0.1 * inch))

    sections = [
        ("Professional Summary", data.get("summary") or data.get("goal", "")),
        ("Certifications", data.get("certifications", "")),
        ("Projects", data.get("projects", "")),
        ("Languages", data.get("languages", "")),
        ("Achievements", data.get("achievements", "")),
        ("Professional Experience", data.get("experience", "")),
        ("Education", data.get("education", "")),
        ("Technical Skills", data.get("skills", "")),
    ]

    for heading, value in sections:
        if not value:
            continue
        story.append(Paragraph(heading, heading_style))
        story.append(Paragraph(value, body_style))
        story.append(Spacer(1, 0.08 * inch))

    doc.build(story)
    buffer.seek(0)
    return buffer

